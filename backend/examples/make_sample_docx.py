"""Build sample_coursework.docx from the «Курсовая» example (examples.json).

    cd backend && python examples/make_sample_docx.py
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
text = next(e["text"] for e in json.loads((HERE.parent / "examples.json").read_text(encoding="utf-8")) if e["id"] == "coursework")
title, *rest = [p.strip() for p in text.split("\n") if p.strip()]

doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(14)
doc.add_heading(title, level=1)
for p in rest:
    if p == "Список литературы":
        doc.add_heading(p, level=2)
    else:
        doc.add_paragraph(p)
doc.save(HERE / "sample_coursework.docx")
print("saved", HERE / "sample_coursework.docx")
